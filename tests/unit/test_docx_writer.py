"""Unit tests for the PDF -> DOCX writer.

Exercises the pure (no-Qt) module against real fixture documents. Atomic-write
side (tempfile + reopen + replace) is also covered here since it does not need
the worker process.
"""

from __future__ import annotations

import re
import zipfile
from pathlib import Path

import pymupdf
import pytest
from docx import Document
from docx.shared import Pt

from fixtures.make_fixtures import (
    ACCENTED,
    BG_PARA,
    LEFT_COL_HEAD,
    PAGE2_LINE,
    PARA_LINE1,
    PARA_LINE2,
    PARA_LINE3,
    PARA_LINE4,
    RIGHT_COL_HEAD,
    RULED_AFTER,
    RULED_TABLE,
    RULED_TITLE,
    TIGHT_ROWS,
    TITLE,
    TRAILING,
)
from openpdfsuite.domain.models import (
    DocxExportOptions,
    FontReference,
    Rect,
    TextLine,
    TextRun,
)
from openpdfsuite.infrastructure.docx.docx_writer import (
    ColumnDetector,
    DocxExportError,
    ImageEmbedder,
    TableDetector,
    write_docx,
    write_docx_atomic,
)

# -- helpers ------------------------------------------------------------------

def _open_pdf(path: Path) -> pymupdf.Document:
    return pymupdf.open(path)


def _all_paragraph_text(doc: Document) -> str:
    return "\n".join(p.text for p in doc.paragraphs)


def _run_formatting(run):
    """Return a dict of the formatting applied to a Run."""
    return {
        "name": run.font.name,
        "size": run.font.size,
        "bold": run.bold,
        "italic": run.italic,
        "color_rgb": (run.font.color.rgb if run.font.color and run.font.color.rgb
                      else None),
    }


# -- basic text ---------------------------------------------------------------

def test_standard_writes_paragraphs(standard_pdf, tmp_path):
    """All known source lines from standard.pdf appear in the DOCX output."""
    pdf = _open_pdf(standard_pdf)
    out = tmp_path / "out.docx"
    result = write_docx(pdf, DocxExportOptions(), out)
    assert result.ok
    assert result.pages_written > 0
    text = _all_paragraph_text(Document(out))
    for expected in (TITLE, PARA_LINE1, PARA_LINE2, PARA_LINE3, PARA_LINE4,
                     TRAILING, PAGE2_LINE):
        assert expected in text, f"missing expected text: {expected!r}"


def test_no_unsupported_items_on_standard(standard_pdf, tmp_path):
    pdf = _open_pdf(standard_pdf)
    out = tmp_path / "out.docx"
    result = write_docx(pdf, DocxExportOptions(), out)
    assert result.ok
    assert result.unsupported_items == [], result.unsupported_items


def test_unicode_text_preserved(unicode_pdf, tmp_path):
    pdf = _open_pdf(unicode_pdf)
    out = tmp_path / "out.docx"
    write_docx(pdf, DocxExportOptions(), out)
    text = _all_paragraph_text(Document(out))
    # PyMuPDF sometimes emits non-breaking spaces for layout reasons; normalize
    # both sides before comparing.
    norm = re.sub(r"\s+", " ", text).strip()
    target = re.sub(r"\s+", " ", ACCENTED).strip()
    assert target in norm, f"missing accented text: {target!r}\ntext was: {norm!r}"


# -- formatting ---------------------------------------------------------------

def test_title_run_has_larger_font(standard_pdf, tmp_path):
    """The 20pt title must come through as a larger size than the 11pt body."""
    pdf = _open_pdf(standard_pdf)
    out = tmp_path / "out.docx"
    write_docx(pdf, DocxExportOptions(), out)
    doc = Document(out)
    # find the title paragraph
    title_para = next(p for p in doc.paragraphs if TITLE in p.text)
    title_size = title_para.runs[0].font.size
    assert title_size == Pt(20)
    # find a body paragraph
    body_para = next(p for p in doc.paragraphs if PARA_LINE1 in p.text)
    body_size = body_para.runs[0].font.size
    assert body_size == Pt(11)
    assert title_size > body_size


def test_white_text_preserves_color(colored_pdf, tmp_path):
    """White text on colored_bg.pdf must carry an RGB color near white."""
    pdf = _open_pdf(colored_pdf)
    out = tmp_path / "out.docx"
    write_docx(pdf, DocxExportOptions(), out)
    doc = Document(out)
    # At least one paragraph contains BG_PARA text and its runs carry near-white.
    found_white = False
    norm_bg = [re.sub(r"\s+", " ", line) for line in BG_PARA]
    for para in doc.paragraphs:
        norm_text = re.sub(r"\s+", " ", para.text)
        if not any(line in norm_text for line in norm_bg):
            continue
        for run in para.runs:
            color_obj = run.font.color
            rgb = color_obj.rgb if (color_obj and color_obj.rgb) else None
            if rgb is None:
                continue
            r, g, b = int(rgb[0]), int(rgb[1]), int(rgb[2])
            if r > 230 and g > 230 and b > 230:
                found_white = True
                break
        if found_white:
            break
    assert found_white, "expected white text on colored_bg.pdf to carry an RGB color"


def test_color_hex_conversion_known_value():
    """Internal: Color -> hex is monotonic and bounded."""
    from openpdfsuite.domain.models import Color
    from openpdfsuite.infrastructure.docx.docx_writer import _hex_color
    assert _hex_color(Color(0, 0, 0)) == "000000"
    assert _hex_color(Color(1, 1, 1)) == "FFFFFF"
    assert _hex_color(Color(0.5, 0.5, 0.5)) == "808080"
    # values clamped
    assert _hex_color(Color(2, -1, 0.5)) == "FF0080"


# -- atomic write -------------------------------------------------------------

def test_atomic_write_creates_file(standard_pdf, tmp_path):
    pdf = _open_pdf(standard_pdf)
    dest = tmp_path / "subdir" / "out.docx"
    dest.parent.mkdir()
    result = write_docx_atomic(pdf, DocxExportOptions(), dest)
    assert result.ok
    assert dest.exists()
    assert Path(result.output_path) == dest
    # ensure the DOCX reopens
    assert len(Document(dest).paragraphs) > 0


def test_atomic_write_keeps_destination_on_failure(standard_pdf, tmp_path):
    """If write_docx raises, the original destination must be untouched."""
    class BoomPdf:
        page_count = 1
        def __iter__(self):
            return iter([])
        def __getitem__(self, i):
            raise RuntimeError("simulated engine failure")

    dest = tmp_path / "out.docx"
    dest.write_bytes(b"ORIGINAL")  # pre-existing content must survive
    with pytest.raises(DocxExportError):
        write_docx_atomic(BoomPdf(), DocxExportOptions(), dest)
    assert dest.read_bytes() == b"ORIGINAL"


def test_atomic_write_rejects_empty_doc(tmp_path):
    """A zero-page PDF must raise a clear DocxExportError, not crash."""
    empty = pymupdf.open()
    out = tmp_path / "out.docx"
    with pytest.raises(DocxExportError, match=r"no pages"):
        write_docx_atomic(empty, DocxExportOptions(), out)


def test_writes_real_docx_extension(standard_pdf, tmp_path):
    pdf = _open_pdf(standard_pdf)
    out = tmp_path / "out.docx"
    write_docx(pdf, DocxExportOptions(), out)
    # DOCX is a ZIP; opening with zipfile confirms the file structure is real.
    assert zipfile.is_zipfile(out)
    with zipfile.ZipFile(out) as z:
        names = set(z.namelist())
        assert "word/document.xml" in names
        # The [Content_Types].xml at the root is the standard DOCX manifest.
        assert "[Content_Types].xml" in names


# -- error surfaces -----------------------------------------------------------

def test_options_default_to_all_enabled():
    """Default options match the v1 plan (embed/tables/columns on)."""
    opts = DocxExportOptions()
    assert opts.embed_images is True
    assert opts.detect_tables is True
    assert opts.detect_columns is True
    assert opts.flow_mode == "formatted"


# -- geometric (ruled) tables ------------------------------------------------

def test_ruled_table_becomes_real_word_table(ruled_table_pdf, tmp_path):
    """A table drawn with vector rulings must export as a Table Grid table."""
    pdf = _open_pdf(ruled_table_pdf)
    out = tmp_path / "out.docx"
    result = write_docx(pdf, DocxExportOptions(), out)
    assert result.ok
    assert result.unsupported_items == []
    doc = Document(out)
    assert len(doc.tables) == 1
    table = doc.tables[0]
    assert len(table.rows) == len(RULED_TABLE)
    assert len(table.columns) == 3
    for ri, expected_row in enumerate(RULED_TABLE):
        actual = [table.cell(ri, ci).text for ci in range(3)]
        assert actual == expected_row, f"row {ri}: {actual}"


def test_ruled_table_does_not_absorb_surrounding_text(ruled_table_pdf, tmp_path):
    """Heading above and paragraph below keep flowing as paragraphs."""
    pdf = _open_pdf(ruled_table_pdf)
    out = tmp_path / "out.docx"
    write_docx(pdf, DocxExportOptions(), out)
    texts = [p.text for p in Document(out).paragraphs if p.text]
    assert RULED_TITLE in texts
    assert RULED_AFTER in texts
    # table cell text must not be duplicated into the flowing paragraphs
    joined = "\n".join(texts)
    assert "Widget" not in joined


def test_colored_background_is_not_a_table(colored_pdf, tmp_path):
    """A filled background rectangle must not be misdetected as a 1x1 table."""
    pdf = _open_pdf(colored_pdf)
    out = tmp_path / "out.docx"
    write_docx(pdf, DocxExportOptions(), out)
    assert Document(out).tables == []


# -- structure regressions ---------------------------------------------------

def test_no_manual_page_breaks(standard_pdf, tmp_path):
    """Sections paginate the document; manual breaks created blank pages."""
    pdf = _open_pdf(standard_pdf)
    out = tmp_path / "out.docx"
    write_docx(pdf, DocxExportOptions(), out)
    with zipfile.ZipFile(out) as z:
        xml = z.read("word/document.xml").decode("utf-8")
    assert 'w:type="page"' not in xml


def test_one_section_per_pdf_page(standard_pdf, tmp_path):
    pdf = _open_pdf(standard_pdf)  # 2 pages
    out = tmp_path / "out.docx"
    result = write_docx(pdf, DocxExportOptions(), out)
    assert result.pages_written == 2  # real page count, not paragraph count
    assert len(Document(out).sections) == 2


def test_lines_merge_into_one_paragraph(standard_pdf, tmp_path):
    """Tightly-leaded PARA_LINE1..4 must merge into a single Word paragraph."""
    pdf = _open_pdf(standard_pdf)
    out = tmp_path / "out.docx"
    write_docx(pdf, DocxExportOptions(), out)
    doc = Document(out)
    merged = [p for p in doc.paragraphs
              if PARA_LINE1 in p.text and PARA_LINE4 in p.text]
    assert merged, "expected the 4 paragraph lines to merge into one paragraph"
    # and the join must read naturally (space-separated, not concatenated)
    assert PARA_LINE1 + " " + PARA_LINE2 in merged[0].text


def test_paragraph_spacing_is_neutralized(standard_pdf, tmp_path):
    """Template paragraph spacing (8pt after, 1.08 lines) must not inflate.

    The trailing line is the last item on page 1, so nothing overrides the
    zeroed spacing and the written values can be asserted directly.
    """
    pdf = _open_pdf(standard_pdf)
    out = tmp_path / "out.docx"
    write_docx(pdf, DocxExportOptions(), out)
    doc = Document(out)
    para = next(p for p in doc.paragraphs if TRAILING in p.text)
    assert para.paragraph_format.space_before.pt == 0
    assert para.paragraph_format.space_after.pt == 0
    assert para.paragraph_format.line_spacing in (None, 1.0)


def test_merged_paragraph_keeps_pdf_leading(standard_pdf, tmp_path):
    """The merged paragraph's line spacing mirrors the PDF's 16/11pt leading."""
    pdf = _open_pdf(standard_pdf)
    out = tmp_path / "out.docx"
    write_docx(pdf, DocxExportOptions(), out)
    doc = Document(out)
    para = next(p for p in doc.paragraphs if PARA_LINE1 in p.text)
    ls = para.paragraph_format.line_spacing
    assert ls and abs(ls - 16 / 11) < 0.05


def test_landscape_page_gets_landscape_section(tmp_path):
    """Page 0 landscape + page 1 portrait: section dims must match exactly."""
    src = pymupdf.open()
    p1 = src.new_page(width=842, height=595)  # landscape
    p1.insert_text((72, 72), "landscape page", fontsize=12, fontname="helv")
    p2 = src.new_page(width=595, height=842)  # portrait
    p2.insert_text((72, 72), "portrait page", fontsize=12, fontname="helv")
    out = tmp_path / "out.docx"
    write_docx(src, DocxExportOptions(), out)
    sections = Document(out).sections
    assert len(sections) == 2
    assert sections[0].page_width == Pt(842)
    assert sections[0].page_height == Pt(595)
    assert sections[1].page_width == Pt(595)
    assert sections[1].page_height == Pt(842)


def test_centered_title_keeps_center_alignment(tmp_path):
    """A centered heading over body text must stay centered in Word."""
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    src = pymupdf.open()
    page = src.new_page(width=595, height=842)
    tw = pymupdf.get_text_length("Centered Heading", fontname="hebo", fontsize=18)
    page.insert_text(((595 - tw) / 2, 100), "Centered Heading",
                     fontsize=18, fontname="hebo")
    page.insert_text((72, 150), PARA_LINE1, fontsize=11, fontname="helv")
    page.insert_text((72, 166), PARA_LINE2, fontsize=11, fontname="helv")
    out = tmp_path / "out.docx"
    write_docx(src, DocxExportOptions(), out)
    doc = Document(out)
    para = next(p for p in doc.paragraphs if "Centered Heading" in p.text)
    assert para.alignment == WD_ALIGN_PARAGRAPH.CENTER


def test_image_placed_inline_with_on_page_size(image_pdf, tmp_path):
    """The image lands inline in reading order at its PDF placement size."""
    pdf = _open_pdf(image_pdf)
    placed = pdf[0].get_image_info(xrefs=True)[0]["bbox"]  # true on-page rect
    out = tmp_path / "out.docx"
    write_docx(pdf, DocxExportOptions(), out)
    doc = Document(out)
    drawings = doc.element.body.findall(
        ".//{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}extent")
    assert drawings
    # EMU: 1pt = 12700 — DOCX width must match the PDF placement width
    w_emu = int(drawings[0].get("cx"))
    placed_w = placed[2] - placed[0]
    assert abs(w_emu - placed_w * 12700) < 0.02 * placed_w * 12700, w_emu
    # text drawn over the image must still be present
    assert any("Text drawn over an image" in p.text for p in doc.paragraphs)


# -- helpers for detector tests ----------------------------------------------

def _make_line(text: str, x0: float, y0: float, x1: float | None = None,
               y1: float | None = None) -> TextLine:
    """Build a minimal TextLine for detector tests."""
    x1 = x1 if x1 is not None else x0 + len(text) * 5.0
    y1 = y1 if y1 is not None else y0 + 12.0
    run = TextRun(text=text, font=FontReference(name="helv", size=11.0),
                  bbox=Rect(x0, y0, x1, y1),
                  origin=type("P", (), {"x": x0, "y": y0 + 10})())
    return TextLine(runs=[run], bbox=Rect(x0, y0, x1, y1), baseline_y=y0 + 10)


# -- column detection --------------------------------------------------------

def test_column_detector_two_bands():
    """Two clear x-clusters should produce two bands."""
    lines = [
        _make_line("Left 1", x0=60, y0=80),
        _make_line("Left 2", x0=60, y0=100),
        _make_line("Left 3", x0=60, y0=120),
        _make_line("Right 1", x0=320, y0=80),
        _make_line("Right 2", x0=320, y0=100),
        _make_line("Right 3", x0=320, y0=120),
    ]
    bands = ColumnDetector().detect(lines)
    assert len(bands) == 2
    assert all(ln.bbox.x0 < 100 for ln in bands[0])
    assert all(ln.bbox.x0 > 300 for ln in bands[1])


def test_column_detector_single_band_when_no_gap():
    """Lines all on the same x should collapse into one band."""
    lines = [_make_line(f"Line {i}", x0=72, y0=80 + i * 12) for i in range(6)]
    bands = ColumnDetector().detect(lines)
    assert len(bands) == 1


def test_columns_pdf_emits_borderless_table(columns_pdf, tmp_path):
    """columns.pdf must be rendered as a Word table with 2 columns."""
    pdf = _open_pdf(columns_pdf)
    out = tmp_path / "out.docx"
    write_docx(pdf, DocxExportOptions(), out)
    doc = Document(out)
    tables = doc.tables
    assert tables, "expected at least one Word table for multi-column page"
    # 2-column layout -> borderless 1-row table
    layout_table = next((t for t in tables if len(t.columns) == 2), None)
    assert layout_table is not None
    # both expected headings land in the cells
    cells_text = " ".join(c.text for c in layout_table._cells)
    assert LEFT_COL_HEAD in cells_text
    assert RIGHT_COL_HEAD in cells_text


# -- table detection ---------------------------------------------------------

def test_table_detector_pipe_rows():
    """Lines with `|` separators form rows; each is split into cells."""
    lines = [_make_line(row, x0=72, y0=80 + i * 12) for i, row in enumerate(TIGHT_ROWS)]
    rows = TableDetector().detect(lines)
    assert rows is not None
    assert len(rows) == 6
    assert all(len(r) == 3 for r in rows), "each TIGHT_ROW must split into 3 cells"
    # cell text content (trailing whitespace before the pipe is fine)
    assert rows[0][0][0].text.strip() == "Row 1 cell one"
    assert rows[0][1][0].text.strip() == "cell two"
    assert rows[0][2][0].text.strip() == "cell three"


def test_table_detector_no_pipes_returns_none():
    """Plain paragraphs without `|` are not tables."""
    lines = [_make_line(f"Line {i}", x0=72, y0=80 + i * 12) for i in range(5)]
    assert TableDetector().detect(lines) is None


def test_tight_pdf_emits_data_table(tight_pdf, tmp_path):
    """tight.pdf must be rendered as a 6-row, 3-col Word table."""
    pdf = _open_pdf(tight_pdf)
    out = tmp_path / "out.docx"
    write_docx(pdf, DocxExportOptions(), out)
    doc = Document(out)
    tables = doc.tables
    data_table = next((t for t in tables if len(t.rows) == 6), None)
    assert data_table is not None, "expected a 6-row Word table"
    assert len(data_table.columns) == 3
    # Spot-check cell content
    cell_text = data_table.cell(0, 0).text + data_table.cell(0, 1).text
    assert "Row 1" in cell_text and "cell one" in cell_text


# -- image embedding ---------------------------------------------------------

def test_image_embedder_extracts_image(image_pdf, tmp_path):
    """image_behind.pdf must produce a DOCX whose package contains an image."""
    pdf = _open_pdf(image_pdf)
    out = tmp_path / "out.docx"
    write_docx(pdf, DocxExportOptions(), out)
    assert zipfile.is_zipfile(out)
    with zipfile.ZipFile(out) as z:
        names = z.namelist()
    image_parts = [n for n in names if n.startswith("word/media/")]
    assert image_parts, f"expected at least one image part; got names={names}"


def test_image_embedder_disabled_does_not_embed(image_pdf, tmp_path):
    """When embed_images=False the DOCX must not contain image parts."""
    pdf = _open_pdf(image_pdf)
    out = tmp_path / "out.docx"
    write_docx(pdf, DocxExportOptions(embed_images=False), out)
    with zipfile.ZipFile(out) as z:
        image_parts = [n for n in z.namelist() if n.startswith("word/media/")]
    assert image_parts == []


def test_image_embedder_class_direct():
    """The class itself can be invoked with a fresh Document."""
    # Build a tiny in-memory PDF with one image
    src = pymupdf.open()
    page = src.new_page(width=200, height=200)
    pix = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 50, 50))
    page.insert_image(pymupdf.Rect(10, 10, 110, 60), pixmap=pix)
    src.save("/tmp/embedder_test.pdf")  # noqa: S108
    try:
        src2 = pymupdf.open("/tmp/embedder_test.pdf")  # noqa: S108
    except Exception:
        src2 = src
    doc = Document()
    embedder = ImageEmbedder()
    n = embedder.embed_for_page(doc, src2[0])
    assert n >= 1
    # DOCX now has at least one paragraph that contains a drawing
    assert any(p._p.findall(".//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}drawing")
               for p in doc.paragraphs)
